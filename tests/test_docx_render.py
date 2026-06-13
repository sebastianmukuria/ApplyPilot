from docx import Document

from applypilot.scoring.docx_render import cover_to_docx, resume_to_docx


RESUME_TEXT = """Jane Doe
jane@example.com | 555-0100 | github.com/janedoe

SUMMARY
Senior engineer focused on reliable backend systems.

TECHNICAL SKILLS
Languages: Python, SQL
Tools: Docker, GitHub Actions

EXPERIENCE
Senior Engineer at Alpha Co
Python | 2022 - Present
- Built a reporting pipeline that cut manual work by 8 hours weekly.
- Improved service latency by 30% with targeted query optimization.

EDUCATION
State University | BS Computer Science
"""


COVER_TEXT = """June 13, 2026

Dear Hiring Manager,

I am excited to apply for the Senior Engineer role. My background is strongest in backend systems, automation, and practical delivery.

Sincerely,
Jane Doe
"""


def test_resume_to_docx_renders_sections_and_bullets(tmp_path):
    path = resume_to_docx(RESUME_TEXT, tmp_path / "resume.docx")

    assert path.exists()
    doc = Document(path)
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    texts = [p.text for p in paragraphs]

    assert texts[0] == "Jane Doe"
    assert "SUMMARY" in texts
    assert "TECHNICAL SKILLS" in texts
    assert "EXPERIENCE" in texts
    assert "EDUCATION" in texts

    bullets = [p for p in paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 2
    assert bullets[0].text.startswith("Built a reporting pipeline")


def test_cover_to_docx_renders_professional_letter(tmp_path):
    path = cover_to_docx(COVER_TEXT, tmp_path / "cover.docx")

    assert path.exists()
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]

    assert texts[0] == "June 13, 2026"
    assert "Dear Hiring Manager," in texts
    assert texts[-1] == "Sincerely,\nJane Doe"


def test_docx_renderers_do_not_crash_on_odd_input(tmp_path):
    assert resume_to_docx("", tmp_path / "empty-resume.docx").exists()
    assert resume_to_docx("Only A Name\nUnstructured line", tmp_path / "odd-resume.docx").exists()
    assert cover_to_docx("", tmp_path / "empty-cover.docx").exists()
