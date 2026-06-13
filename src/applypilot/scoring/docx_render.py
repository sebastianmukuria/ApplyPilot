"""DOCX rendering for tailored resumes and cover letters."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT_NAME = "Georgia"

RESUME_MARGIN_IN = 0.7
RESUME_NAME_PT = 16
RESUME_CONTACT_PT = 9.5
RESUME_SECTION_PT = 10.5
RESUME_BODY_PT = 10
RESUME_PARAGRAPH_AFTER_PT = 3
RESUME_SECTION_BEFORE_PT = 7
RESUME_BULLET_LEFT_IN = 0.22
RESUME_BULLET_HANGING_IN = 0.14

COVER_MARGIN_IN = 1.0
COVER_BODY_PT = 11
COVER_PARAGRAPH_AFTER_PT = 9
COVER_LINE_SPACING = 1.15

SECTION_BORDER_COLOR = "888888"
SECTION_BORDER_SIZE = "4"

KNOWN_HEADERS = {
    "SUMMARY",
    "TECHNICAL SKILLS",
    "EXPERIENCE",
    "EDUCATION",
    "PROJECTS",
    "SKILLS",
    "CERTIFICATIONS",
}


def resume_to_docx(text: str, out_path: Path) -> Path:
    """Render structured resume text to a DOCX file."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _set_margins(document, RESUME_MARGIN_IN)
    _set_default_font(document, RESUME_BODY_PT)

    lines = _trim_blank_edges(text.splitlines())
    header_end = _first_header_index(lines)
    header_lines = [line.strip() for line in lines[:header_end] if line.strip()]
    body_lines = lines[header_end:] if header_end < len(lines) else []

    if header_lines:
        name = document.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name.paragraph_format.space_after = Pt(2)
        run = name.add_run(header_lines[0])
        _format_run(run, RESUME_NAME_PT, bold=True)
        run.font.small_caps = True

        for line in header_lines[1:]:
            contact = document.add_paragraph()
            contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact.paragraph_format.space_after = Pt(3)
            _format_run(contact.add_run(line), RESUME_CONTACT_PT)
    elif not body_lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)

    for raw_line in body_lines:
        stripped = raw_line.strip()
        if not stripped:
            continue
        upper = stripped.upper()
        if upper in KNOWN_HEADERS:
            _add_section_header(document, upper)
        elif stripped.startswith("- "):
            _add_bullet(document, stripped[2:].strip())
        elif stripped.startswith("\u2022 "):
            _add_bullet(document, stripped[2:].strip())
        else:
            paragraph = document.add_paragraph()
            _tight_paragraph(paragraph, RESUME_PARAGRAPH_AFTER_PT)
            _format_run(paragraph.add_run(stripped), RESUME_BODY_PT)

    document.save(out_path)
    return out_path


def cover_to_docx(text: str, out_path: Path) -> Path:
    """Render cover-letter text to a DOCX file."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _set_margins(document, COVER_MARGIN_IN)
    _set_default_font(document, COVER_BODY_PT)

    paragraphs = _cover_paragraphs(text)
    if not paragraphs:
        paragraph = document.add_paragraph()
        _cover_paragraph_format(paragraph)
    for idx, paragraph_text in enumerate(paragraphs):
        paragraph = document.add_paragraph()
        _cover_paragraph_format(paragraph)
        if idx == 0 and _looks_like_date(paragraph_text):
            paragraph.paragraph_format.space_after = Pt(14)
        elif idx == len(paragraphs) - 1 and _looks_like_signature(paragraph_text):
            paragraph.paragraph_format.space_before = Pt(8)
        _add_text_with_breaks(paragraph, paragraph_text, COVER_BODY_PT)

    document.save(out_path)
    return out_path


def _set_margins(document: Document, margin_in: float) -> None:
    margin = Inches(margin_in)
    for section in document.sections:
        section.top_margin = margin
        section.right_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin


def _set_default_font(document: Document, size_pt: float) -> None:
    style = document.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(size_pt)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)


def _format_run(run, size_pt: float, *, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size_pt)
    run.bold = bold


def _tight_paragraph(paragraph, after_pt: float) -> None:
    paragraph.paragraph_format.space_after = Pt(after_pt)
    paragraph.paragraph_format.line_spacing = 1.0


def _add_section_header(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(RESUME_SECTION_BEFORE_PT)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    _format_run(run, RESUME_SECTION_PT, bold=True)
    _add_bottom_border(paragraph)


def _add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    _tight_paragraph(paragraph, 2)
    paragraph.paragraph_format.left_indent = Inches(RESUME_BULLET_LEFT_IN)
    paragraph.paragraph_format.first_line_indent = -Inches(RESUME_BULLET_HANGING_IN)
    _format_run(paragraph.add_run(text), RESUME_BODY_PT)


def _add_bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), SECTION_BORDER_SIZE)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), SECTION_BORDER_COLOR)
    p_bdr.append(bottom)


def _first_header_index(lines: list[str]) -> int:
    for idx, line in enumerate(lines):
        if line.strip().upper() in KNOWN_HEADERS:
            return idx
    return len(lines)


def _trim_blank_edges(lines: list[str]) -> list[str]:
    start = 0
    end = len(lines)
    while start < end and not lines[start].strip():
        start += 1
    while end > start and not lines[end - 1].strip():
        end -= 1
    return [line.rstrip() for line in lines[start:end]]


def _cover_paragraphs(text: str) -> list[str]:
    stripped = text.strip()
    if not stripped:
        return []
    return [p.strip() for p in re.split(r"\n\s*\n", stripped) if p.strip()]


def _cover_paragraph_format(paragraph) -> None:
    paragraph.paragraph_format.space_after = Pt(COVER_PARAGRAPH_AFTER_PT)
    paragraph.paragraph_format.line_spacing = COVER_LINE_SPACING


def _add_text_with_breaks(paragraph, text: str, size_pt: float) -> None:
    parts = text.splitlines() or [text]
    for idx, part in enumerate(parts):
        if idx:
            paragraph.add_run().add_break()
        _format_run(paragraph.add_run(part.strip()), size_pt)


def _looks_like_date(text: str) -> bool:
    return bool(re.search(r"\b\d{4}\b", text)) or bool(
        re.match(
            r"^(January|February|March|April|May|June|July|August|September|October|November|December)\b",
            text,
        )
    )


def _looks_like_signature(text: str) -> bool:
    return bool(re.match(r"^(Sincerely|Best|Regards|Thank you|Thanks)\b", text.strip(), re.I))
